"""Unit tests for ExcelConnector: unmerging, hidden rows, totals dropping."""

from pathlib import Path

import openpyxl
import pytest
from fieldkit_ingest.connectors.base import IngestContext
from fieldkit_ingest.connectors.excel_connector import ExcelConnector


@pytest.mark.asyncio
async def test_excel_fixture_unmerging_and_filtering(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "test_planning.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Planung"

    # Row 1-2: Title banners
    ws["A1"] = "HALDEN INDUSTRIE GMBH - PLANUNG"
    ws["A2"] = "VERTRAULICH"

    # Row 3: Merged Category headers
    ws.merge_cells("A3:B3")
    ws["A3"] = "Lieferantendaten"
    ws.merge_cells("C3:D3")
    ws["C3"] = "Konditionen"

    # Row 4: Column headers
    ws["A4"] = "Lieferanten-Nr"
    ws["B4"] = "Firmenname"
    ws["C4"] = "Ort"
    ws["D4"] = "Planvolumen (EUR)"

    # Row 5: Visible data row 1
    ws["A5"] = "HAM-001"
    ws["B5"] = "Nordic Valve GmbH"
    ws["C5"] = "Hamburg"
    ws["D5"] = 15000.0

    # Row 6: Hidden data row (must be skipped)
    ws["A6"] = "HAM-002"
    ws["B6"] = "Deleted Vendor KG"
    ws["C6"] = "Hamburg"
    ws["D6"] = 5000.0
    ws.row_dimensions[6].hidden = True

    # Row 7: Visible data row 2
    ws["A7"] = "HAM-003"
    ws["B7"] = "Baltic Seals AG"
    ws["C7"] = "Rostock"
    ws["D7"] = 25000.0

    # Row 8: Totals row (must be dropped)
    ws["A8"] = "GESAMTVOLUMEN"
    ws["D8"] = "=SUM(D5:D7)"

    wb.save(xlsx_file)

    connector = ExcelConnector(xlsx_file)
    ctx = IngestContext(run_id="test_run", source_name="test_excel")

    batches = [b async for b in connector.extract(ctx)]
    assert len(batches) == 1
    records = batches[0].records

    # Should only contain HAM-001 and HAM-003 (hidden HAM-002 and totals row dropped)
    assert len(records) == 2
    assert records[0]["Lieferanten-Nr"] == "HAM-001"
    assert records[0]["Firmenname"] == "Nordic Valve GmbH"
    assert records[1]["Lieferanten-Nr"] == "HAM-003"
    assert records[1]["Firmenname"] == "Baltic Seals AG"
