"""Generate synthetic Halden Industrie ERP datasets and evaluation ground truth.

Features:
- Samples 3,000 real European companies from data/ref/gleif_eu.parquet
- Generates multi-source vendor records with deterministic corruptions:
  * data/raw/erp/erp_vendors.csv (Latin-1, semicolon delimiter, ~3,400 rows)
  * data/raw/excel/planning_de_ham.xlsx (1,200 rows, merged headers, title rows, totals)
  * data/raw/excel/planning_de_str.xlsx (950 rows, merged headers, title rows, totals)
  * data/raw/soap/soap_seed.json (800 mock SOAP certificates)
- Generates Halden operational manufacturing data:
  * data/raw/erp/parts.csv (500 parts)
  * data/raw/erp/purchase_orders.csv (50,000 POs with mixed date formats & comma decimals)
  * data/raw/erp/quality_incidents.csv (3,000 quality incidents)
- Emits ground truth mapping:
  * data/truth/entity_mapping.parquet (record_id, source, true_entity_id, lei)
- Emits evaluation benchmark dataset:
  * evals/p1/gold_pairs.parquet (1,000 true match pairs + 1,000 hard negative pairs)
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sys
import time
from pathlib import Path
from typing import Any

# Ensure workspace root is in sys.path
ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import duckdb  # noqa: E402
import openpyxl  # noqa: E402
import pyarrow as pa  # noqa: E402
import pyarrow.parquet as pq  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from scripts.corruptions import apply_vendor_corruptions  # noqa: E402

# Directory paths
DATA_REF = Path("data/ref")
DATA_RAW_ERP = Path("data/raw/erp")
DATA_RAW_EXCEL = Path("data/raw/excel")
DATA_RAW_SOAP = Path("data/raw/soap")
DATA_TRUTH = Path("data/truth")
EVALS_P1 = Path("evals/p1")


def ensure_directories() -> None:
    for d in [DATA_RAW_ERP, DATA_RAW_EXCEL, DATA_RAW_SOAP, DATA_TRUTH, EVALS_P1]:
        d.mkdir(parents=True, exist_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate synthetic Halden Industrie ERP datasets."
    )
    parser.add_argument("--seed", type=int, default=42, help="Random seed for reproducibility")
    parser.add_argument(
        "--n-entities", type=int, default=3000, help="Number of seed entities to sample"
    )
    return parser.parse_args()


def load_seed_entities(n_entities: int, seed: int) -> list[dict[str, Any]]:
    parquet_path = DATA_REF / "gleif_eu.parquet"
    if not parquet_path.exists():
        raise FileNotFoundError(
            f"Reference dataset not found at {parquet_path}. Run scripts/load_gleif.py first."
        )

    con = duckdb.connect()
    query = f"""
    SELECT
        lei,
        legal_name,
        legal_form,
        street,
        city,
        postcode,
        country
    FROM '{parquet_path.as_posix()}'
    WHERE legal_name IS NOT NULL AND city != ''
    ORDER BY hash(concat(lei, '{seed}'))
    LIMIT {n_entities};
    """
    rows = con.execute(query).fetchall()
    cols = ["lei", "legal_name", "legal_form", "street", "city", "postcode", "country"]

    entities = []
    for idx, r in enumerate(rows, 1):
        d = dict(zip(cols, r, strict=False))
        d["entity_id"] = f"ENT-{idx:05d}"
        entities.append(d)

    return entities


def generate_erp_vendors(
    entities: list[dict[str, Any]],
    rng: random.Random,
    mapping_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Generates 3,400 ERP vendor records with Latin-1 encoding and semicolon delimiter."""
    out_path = DATA_RAW_ERP / "erp_vendors.csv"
    erp_records = []

    # Take all 3,000 entities + duplicate 400 with corruption
    pool = list(entities)
    duplicates = rng.choices(entities, k=400)
    all_targets = pool + duplicates
    rng.shuffle(all_targets)

    terms_choices = ["Net 30", "Net 60", "Net 14", "2% 10 Net 30", "Immediate"]

    for idx, target in enumerate(all_targets, 1):
        record_id = f"ERP-{idx:05d}"
        corrupted = apply_vendor_corruptions(target, rng)

        currency = (
            "EUR"
            if target["country"] in ("DE", "AT")
            else ("PLN" if target["country"] == "PL" else "CZK")
        )
        row = {
            "vendor_id": record_id,
            "company_name": corrupted["legal_name"],
            "street": corrupted["street"],
            "city": corrupted["city"],
            "postal_code": corrupted["postcode"],
            "country": corrupted["country"],
            "payment_terms": rng.choice(terms_choices),
            "currency": currency,
        }
        erp_records.append(row)

        mapping_records.append(
            {
                "record_id": record_id,
                "source": "erp_vendors",
                "true_entity_id": target["entity_id"],
                "lei": target["lei"],
                "company_name": corrupted["legal_name"],
                "street": corrupted["street"],
                "city": corrupted["city"],
                "postal_code": corrupted["postcode"],
                "country": corrupted["country"],
            }
        )

    # Write Latin-1 semicolon-separated CSV
    with open(out_path, mode="w", encoding="latin-1", errors="replace", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "vendor_id",
                "company_name",
                "street",
                "city",
                "postal_code",
                "country",
                "payment_terms",
                "currency",
            ],
            delimiter=";",
        )
        writer.writeheader()
        writer.writerows(erp_records)

    return erp_records


def generate_excel_planning(
    filename: str,
    title: str,
    n_rows: int,
    id_prefix: str,
    entities: list[dict[str, Any]],
    rng: random.Random,
    mapping_records: list[dict[str, Any]],
) -> None:
    """Generates realistic enterprise planning spreadsheets with merged headers and formatting."""
    out_path = DATA_RAW_EXCEL / filename
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lieferanten_Planung"

    # Banner Title Rows
    ws.merge_cells("A1:G1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = "INTERNES ARBEITSDOKUMENT -- HALDEN INDUSTRIE GMBH -- STAND 2026"
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Merged Category Headers (Row 3 & 4)
    ws.merge_cells("A3:E3")
    ws["A3"] = "Stammdaten Lieferant"
    ws["A3"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws["A3"].fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    ws["A3"].alignment = Alignment(horizontal="center")

    ws.merge_cells("F3:G3")
    ws["F3"] = "Planungskonditionen"
    ws["F3"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    ws["F3"].fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
    ws["F3"].alignment = Alignment(horizontal="center")

    headers = [
        "Lieferanten-Nr",
        "Firmenname",
        "Strasse / Hausnummer",
        "Ort",
        "PLZ / Land",
        "Kategorie",
        "Planvolumen (EUR)",
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    categories = [
        "Hydraulikkomponenten",
        "Pneumatik",
        "Dichtungstechnik",
        "Guss- und Schmiedeteile",
        "Rohrverbindungen",
        "Sensorik",
    ]

    # Sample rows from entities
    sampled = rng.choices(entities, k=n_rows)
    current_row = 5

    for idx, target in enumerate(sampled, 1):
        record_id = f"{id_prefix}-{idx:04d}"
        corrupted = apply_vendor_corruptions(target, rng)

        plz_land = f"{corrupted['postcode']} {corrupted['country']}".strip()
        vol = round(rng.uniform(5000.0, 500000.0), 2)

        ws.cell(row=current_row, column=1, value=record_id)
        ws.cell(row=current_row, column=2, value=corrupted["legal_name"])
        ws.cell(row=current_row, column=3, value=corrupted["street"])
        ws.cell(row=current_row, column=4, value=corrupted["city"])
        ws.cell(row=current_row, column=5, value=plz_land)
        ws.cell(row=current_row, column=6, value=rng.choice(categories))
        ws.cell(row=current_row, column=7, value=vol)

        # Simulate a few hidden rows (e.g. 2% probability)
        if rng.random() < 0.02:
            ws.row_dimensions[current_row].hidden = True

        mapping_records.append(
            {
                "record_id": record_id,
                "source": filename,
                "true_entity_id": target["entity_id"],
                "lei": target["lei"],
                "company_name": corrupted["legal_name"],
                "street": corrupted["street"],
                "city": corrupted["city"],
                "postal_code": corrupted["postcode"],
                "country": corrupted["country"],
            }
        )
        current_row += 1

    # Totals Row at the bottom
    ws.cell(row=current_row, column=1, value="GESAMTVOLUMEN")
    ws.cell(row=current_row, column=1).font = Font(name="Arial", size=10, bold=True)
    ws.cell(row=current_row, column=7, value=f"=SUM(G5:G{current_row - 1})")
    ws.cell(row=current_row, column=7).font = Font(name="Arial", size=10, bold=True)

    wb.save(out_path)


def generate_soap_seed(
    entities: list[dict[str, Any]],
    rng: random.Random,
    mapping_records: list[dict[str, Any]],
) -> None:
    """Generates 800 vendor certificates in JSON format for the legacy SOAP API."""
    out_path = DATA_RAW_SOAP / "soap_seed.json"
    sampled = rng.choices(entities, k=800)
    soap_records = []

    for idx, target in enumerate(sampled, 1):
        record_id = f"SOAP-{idx:04d}"
        corrupted = apply_vendor_corruptions(target, rng)

        cert = {
            "soap_vendor_id": record_id,
            "company_name": corrupted["legal_name"],
            "address": {
                "street": corrupted["street"],
                "city": corrupted["city"],
                "postcode": corrupted["postcode"],
                "country": corrupted["country"],
            },
            "certifications": [
                {
                    "standard": "ISO 9001:2015",
                    "certificate_id": f"CERT-{rng.randint(100000, 999999)}",
                    "issued_date": f"202{rng.randint(0, 3)}-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}",
                    "expiry_date": f"202{rng.randint(4, 7)}-{rng.randint(1, 12):02d}-{rng.randint(1, 28):02d}",
                    "audit_status": rng.choice(["COMPLIANT", "RENEWAL_PENDING", "ACTIVE"]),
                }
            ],
        }
        soap_records.append(cert)

        mapping_records.append(
            {
                "record_id": record_id,
                "source": "soap_seed",
                "true_entity_id": target["entity_id"],
                "lei": target["lei"],
                "company_name": corrupted["legal_name"],
                "street": corrupted["street"],
                "city": corrupted["city"],
                "postal_code": corrupted["postcode"],
                "country": corrupted["country"],
            }
        )

    with open(out_path, mode="w", encoding="utf-8") as f:
        json.dump(soap_records, f, indent=2, ensure_ascii=False)


def generate_parts_and_orders(
    erp_vendors: list[dict[str, Any]],
    rng: random.Random,
) -> list[str]:
    """Generates 500 parts, 50,000 purchase orders, and 3,000 quality incidents."""
    # 1. 500 Parts
    parts_path = DATA_RAW_ERP / "parts.csv"
    categories = [
        "Hydraulikventile",
        "Hydraulikzylinder",
        "Dichtungssatz",
        "Hochdruckpumpen",
        "Hydraulikschlaeuche",
        "Drucksensoren",
    ]
    parts = []
    part_numbers = []
    for i in range(1, 501):
        p_num = f"HX-{100 + (i % 500):03d}-{i:03d}"
        cat = rng.choice(categories)
        cost = round(rng.uniform(15.0, 1850.0), 2)
        parts.append(
            {
                "part_number": p_num,
                "name": f"Halden {cat} Typ {i:03d}",
                "category": cat,
                "cost_eur": cost,
            }
        )
        part_numbers.append(p_num)

    with open(parts_path, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["part_number", "name", "category", "cost_eur"])
        writer.writeheader()
        writer.writerows(parts)

    # 2. 50,000 Purchase Orders
    po_path = DATA_RAW_ERP / "purchase_orders.csv"
    vendor_ids = [v["vendor_id"] for v in erp_vendors]
    statuses = ["DELIVERED", "DELIVERED", "DELIVERED", "OPEN", "CANCELLED", "PENDING_INSPECTION"]

    po_records = []
    po_numbers = []

    for i in range(1, 50001):
        po_num = f"PO-{i:06d}"
        po_numbers.append(po_num)
        v_id = rng.choice(vendor_ids)
        p_num = rng.choice(part_numbers)
        qty = rng.randint(5, 500)
        unit_price = round(rng.uniform(20.0, 2400.0), 2)

        # Comma decimals (German format: 1250,50)
        price_str = f"{unit_price:.2f}".replace(".", ",")

        # Mixed date formats: DD.MM.YYYY, YYYY-MM-DD, MM/DD/YY
        day = rng.randint(1, 28)
        month = rng.randint(1, 12)
        year = rng.randint(2023, 2026)
        fmt = rng.choice(["de", "iso", "us"])
        if fmt == "de":
            date_str = f"{day:02d}.{month:02d}.{year}"
        elif fmt == "iso":
            date_str = f"{year}-{month:02d}-{day:02d}"
        else:
            date_str = f"{month:02d}/{day:02d}/{year % 100:02d}"

        po_records.append(
            {
                "po_number": po_num,
                "vendor_id": v_id,
                "part_number": p_num,
                "quantity": qty,
                "unit_price": price_str,
                "order_date": date_str,
                "status": rng.choice(statuses),
            }
        )

    with open(po_path, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "po_number",
                "vendor_id",
                "part_number",
                "quantity",
                "unit_price",
                "order_date",
                "status",
            ],
        )
        writer.writeheader()
        writer.writerows(po_records)

    # 3. 3,000 Quality Incidents
    incidents_path = DATA_RAW_ERP / "quality_incidents.csv"
    defect_types = [
        "DIMENSIONAL_OUT_OF_SPEC",
        "SURFACE_DEFECT",
        "LEAKAGE_PRESSURE_TEST",
        "MATERIAL_FATIGUE",
        "CONTAMINATION",
        "MISSING_CERTIFICATE",
    ]
    severities = ["LOW", "MEDIUM", "HIGH", "CRITICAL"]

    incident_notes = [
        (
            "Dichtflaeche weist Riefen auf, Drucktest bei 350 bar fehlgeschlagen.",
            "Sealing surface has scratches, 350 bar pressure test failed.",
        ),
        (
            "Bohrungsmass um 0.12mm ueberschritten. Nacharbeit erforderlich.",
            "Bore dimension exceeded by 0.12mm. Rework required.",
        ),
        (
            "Materialzertifikat 3.1 nach EN 10204 fehlt in Lieferpapieren.",
            "Inspection certificate 3.1 per EN 10204 missing from papers.",
        ),
        (
            "Spannungsriss im Zylindergehaeuse unter Lastzyklus festgestellt.",
            "Stress crack detected in cylinder housing under load cycle.",
        ),
        (
            "Oelverschmutzung in Innenkammer ueberschreitet Grenzwert.",
            "Oil contamination in internal cavity exceeds threshold.",
        ),
        (
            "Gewinde unzulaessig deformiert, Montagesperre verhaengt.",
            "Thread deformed beyond tolerance, assembly hold initiated.",
        ),
    ]

    incidents = []
    sampled_pos = rng.sample(po_numbers, 3000)
    for idx, po in enumerate(sampled_pos, 1):
        defect = rng.choice(defect_types)
        sev = rng.choice(severities)
        notes_de, notes_en = rng.choice(incident_notes)
        incidents.append(
            {
                "incident_id": f"INC-{idx:05d}",
                "po_number": po,
                "defect_type": defect,
                "severity": sev,
                "notes_de": notes_de,
                "notes_en": notes_en,
            }
        )

    with open(incidents_path, mode="w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "incident_id",
                "po_number",
                "defect_type",
                "severity",
                "notes_de",
                "notes_en",
            ],
        )
        writer.writeheader()
        writer.writerows(incidents)

    return po_numbers


def generate_evaluation_gold_pairs(
    mapping_records: list[dict[str, Any]],
    rng: random.Random,
) -> None:
    """Generates 1,000 true match pairs and 1,000 hard negative pairs for Entity Resolution evals."""
    gold_pairs_path = EVALS_P1 / "gold_pairs.parquet"

    # Group records by true_entity_id
    by_entity: dict[str, list[dict[str, Any]]] = {}
    for r in mapping_records:
        by_entity.setdefault(r["true_entity_id"], []).append(r)

    # 1. Select 1,000 True Match pairs across sources
    true_pairs = []
    candidates = [recs for recs in by_entity.values() if len(recs) >= 2]
    rng.shuffle(candidates)

    for recs in candidates:
        if len(true_pairs) >= 1000:
            break
        a, b = rng.sample(recs, 2)
        true_pairs.append(
            {
                "id_a": a["record_id"],
                "source_a": a["source"],
                "name_a": a["company_name"],
                "street_a": a["street"],
                "city_a": a["city"],
                "postcode_a": a["postal_code"],
                "country_a": a["country"],
                "id_b": b["record_id"],
                "source_b": b["source"],
                "name_b": b["company_name"],
                "street_b": b["street"],
                "city_b": b["city"],
                "postcode_b": b["postal_code"],
                "country_b": b["country"],
                "label": 1,
            }
        )

    # If candidates wasn't enough, cycle until exactly 1,000
    while len(true_pairs) < 1000:
        recs = rng.choice(candidates)
        a, b = rng.sample(recs, 2)
        true_pairs.append(
            {
                "id_a": a["record_id"],
                "source_a": a["source"],
                "name_a": a["company_name"],
                "street_a": a["street"],
                "city_a": a["city"],
                "postcode_a": a["postal_code"],
                "country_a": a["country"],
                "id_b": b["record_id"],
                "source_b": b["source"],
                "name_b": b["company_name"],
                "street_b": b["street"],
                "city_b": b["city"],
                "postcode_b": b["postal_code"],
                "country_b": b["country"],
                "label": 1,
            }
        )

    # 2. Select 1,000 Hard Negative pairs using DuckDB Jaro-Winkler
    con = duckdb.connect()
    records_table = pa.Table.from_pylist(mapping_records)
    con.register("records", records_table)

    query = """
    SELECT
        a.record_id AS id_a,
        a.source AS source_a,
        a.company_name AS name_a,
        a.street AS street_a,
        a.city AS city_a,
        a.postal_code AS postcode_a,
        a.country AS country_a,
        b.record_id AS id_b,
        b.source AS source_b,
        b.company_name AS name_b,
        b.street AS street_b,
        b.city AS city_b,
        b.postal_code AS postcode_b,
        b.country AS country_b,
        0 AS label
    FROM records a
    JOIN records b
      ON a.city = b.city
     AND a.true_entity_id < b.true_entity_id
    WHERE jaro_winkler_similarity(a.company_name, b.company_name) > 0.78
    LIMIT 2000;
    """
    hard_neg_rows = con.execute(query).fetchall()
    cols = [
        "id_a",
        "source_a",
        "name_a",
        "street_a",
        "city_a",
        "postcode_a",
        "country_a",
        "id_b",
        "source_b",
        "name_b",
        "street_b",
        "city_b",
        "postcode_b",
        "country_b",
        "label",
    ]

    hard_neg_pairs = [dict(zip(cols, r, strict=False)) for r in hard_neg_rows]

    # If same city gave fewer than 1000 pairs, loosen to same country with higher threshold
    if len(hard_neg_pairs) < 1000:
        needed = 1000 - len(hard_neg_pairs)
        backup_query = f"""
        SELECT
            a.record_id AS id_a,
            a.source AS source_a,
            a.company_name AS name_a,
            a.street AS street_a,
            a.city AS city_a,
            a.postal_code AS postcode_a,
            a.country AS country_a,
            b.record_id AS id_b,
            b.source AS source_b,
            b.company_name AS name_b,
            b.street AS street_b,
            b.city AS city_b,
            b.postal_code AS postcode_b,
            b.country AS country_b,
            0 AS label
        FROM records a
        JOIN records b
          ON a.country = b.country
         AND a.true_entity_id < b.true_entity_id
        WHERE jaro_winkler_similarity(a.company_name, b.company_name) > 0.85
        LIMIT {needed + 100};
        """
        extra_rows = con.execute(backup_query).fetchall()
        hard_neg_pairs.extend([dict(zip(cols, r, strict=False)) for r in extra_rows])

    hard_neg_pairs = rng.sample(hard_neg_pairs, 1000)

    # Combine exactly 1,000 true + 1,000 negative = 2,000 gold pairs
    all_gold_pairs = true_pairs[:1000] + hard_neg_pairs[:1000]
    rng.shuffle(all_gold_pairs)

    table = pa.Table.from_pylist(all_gold_pairs)
    pq.write_table(table, gold_pairs_path, compression="zstd")

    print(f"\n[EVALS] Emitted {len(all_gold_pairs):,} gold pairs to {gold_pairs_path}")
    print(f"  - True matches (label=1): {sum(1 for p in all_gold_pairs if p['label'] == 1):,}")
    print(f"  - Hard negatives (label=0): {sum(1 for p in all_gold_pairs if p['label'] == 0):,}")
    print(f"  - Parquet size: {gold_pairs_path.stat().st_size / (1024 * 1024):.2f} MB")


def main() -> None:
    args = parse_args()
    start_time = time.perf_counter()

    print(f"=== Generating Halden Industrie ERP Datasets (Seed: {args.seed}) ===")
    ensure_directories()
    rng = random.Random(args.seed)

    # 1. Sample Clean Seed Entities from GLEIF
    print(f"1. Sampling {args.n_entities:,} European entities from data/ref/gleif_eu.parquet...")
    entities = load_seed_entities(args.n_entities, args.seed)
    print(f"   [+] Sampled {len(entities):,} entities.")

    # Storage for ground truth mapping
    mapping_records: list[dict[str, Any]] = []

    # 2. Generate ERP Vendors
    print("2. Generating data/raw/erp/erp_vendors.csv (Latin-1, semicolon delimiter)...")
    erp_vendors = generate_erp_vendors(entities, rng, mapping_records)
    print(f"   [+] Generated {len(erp_vendors):,} ERP vendor rows.")

    # 3. Generate Planning Spreadsheets
    print("3. Generating messy planning spreadsheets in data/raw/excel/...")
    generate_excel_planning(
        "planning_de_ham.xlsx",
        "HALDEN INDUSTRIE GMBH -- WERK HAMBURG -- BESCHAFFUNGSPLANUNG 2026",
        1200,
        "HAM",
        entities,
        rng,
        mapping_records,
    )
    print("   [+] Generated planning_de_ham.xlsx (1,200 rows with merged headers & formulas).")

    generate_excel_planning(
        "planning_de_str.xlsx",
        "HALDEN INDUSTRIE GMBH -- WERK STRALSUND -- BESCHAFFUNGSPLANUNG 2026",
        950,
        "STR",
        entities,
        rng,
        mapping_records,
    )
    print("   [+] Generated planning_de_str.xlsx (950 rows with merged headers & formulas).")

    # 4. Generate SOAP API Seed
    print("4. Generating legacy SOAP API certificate seed in data/raw/soap/soap_seed.json...")
    generate_soap_seed(entities, rng, mapping_records)
    print("   [+] Generated 800 SOAP vendor certificate records.")

    # 5. Generate Operational Tables (Parts, POs, Incidents)
    print("5. Generating manufacturing operations (500 parts, 50,000 POs, 3,000 incidents)...")
    generate_parts_and_orders(erp_vendors, rng)
    print("   [+] Parts, mixed-format purchase orders, and bilingual quality incidents written.")

    # 6. Save Ground Truth Mapping
    truth_path = DATA_TRUTH / "entity_mapping.parquet"
    truth_cols = ["record_id", "source", "true_entity_id", "lei"]
    truth_data = [{k: r[k] for k in truth_cols} for r in mapping_records]
    truth_table = pa.Table.from_pylist(truth_data)
    pq.write_table(truth_table, truth_path, compression="zstd")
    print(f"6. Saved ground truth mapping ({len(truth_data):,} records) to {truth_path}")

    # 7. Generate Evaluation Gold Pairs
    print(
        "7. Generating 2,000 gold benchmark pairs for Entity Resolution (evals/p1/gold_pairs.parquet)..."
    )
    generate_evaluation_gold_pairs(mapping_records, rng)

    elapsed = time.perf_counter() - start_time
    print(f"\n[SUCCESS] All Halden ERP datasets generated in {elapsed:.2f} seconds.")


if __name__ == "__main__":
    main()
